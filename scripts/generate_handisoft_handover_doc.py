#!/usr/bin/env python3
"""Generate Handisoft CRM Data Handover Word document."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor
from datetime import date


def set_cell_shading(cell, color_hex: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color_hex)
    shading.set(qn("w:val"), "clear")
    cell._tc.get_or_add_tcPr().append(shading)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(11)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(text, style="List Bullet")
    if level > 0:
        p.paragraph_format.left_indent = Inches(0.25 * level)


def add_numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], header_color: str = "1F4E79") -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        hdr_cells[i].text = header
        set_cell_shading(hdr_cells[i], header_color)
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(10)

    for row_idx, row_data in enumerate(rows):
        row_cells = table.rows[row_idx + 1].cells
        for col_idx, cell_text in enumerate(row_data):
            row_cells[col_idx].text = cell_text
            for paragraph in row_cells[col_idx].paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph()


def add_checkbox_section(doc: Document, title: str, items: list[str]) -> None:
    add_heading(doc, title, level=3)
    for item in items:
        p = doc.add_paragraph()
        p.add_run("☐  ").font.size = Pt(11)
        p.add_run(item).font.size = Pt(11)


def build_document() -> Document:
    doc = Document()

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Handisoft CRM Data Extraction\nHandover Guide")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("For Dummies Edition — Business Sale / Vendor Handover")
    sub_run.font.size = Pt(14)
    sub_run.italic = True

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(f"Prepared: {date.today().strftime('%d %B %Y')}\n")
    meta_run.font.size = Pt(11)
    meta.add_run("Systems: HandiTax · Handisoft Practice Manager (Contact / Jobflow) · Document Manager").font.size = Pt(11)

    doc.add_page_break()

    # Table of contents placeholder
    add_heading(doc, "Contents", 1)
    toc_items = [
        "1. Executive Summary",
        "2. Before You Start",
        "3. Where Your Data Lives",
        "4. Practice Manager Contact (CRM) Export",
        "5. HandiTax Client Export",
        "6. Jobs, WIP & Billing",
        "7. Full Server Backup Procedure",
        "8. Searching W:\\ABC_Data for Manuals",
        "9. Known Gaps & Workarounds",
        "10. Troubleshooting",
        "11. Legal & Privacy Checklist (Australia)",
        "12. 48-Hour Run Sheet",
        "13. Useful Links & Contacts",
        "14. Sign-Off Page",
    ]
    for item in toc_items:
        add_bullet(doc, item)

    doc.add_page_break()

    # Section 1
    add_heading(doc, "1. Executive Summary", 1)
    add_para(
        doc,
        "There is no single \"Export Everything\" button in Handisoft. To complete a business "
        "handover when the vendor will not assist, you need three things:",
    )
    add_numbered(doc, "Human-readable exports (Excel/CSV) via built-in reports — for migration to new software.")
    add_numbered(doc, "Full server backups (HSoft folder + SQL database) — for completeness and disaster recovery.")
    add_numbered(doc, "Document folder copies (DocBase + ABC_Data) — for PDFs, letters, scans, and templates.")
    add_para(doc, "Do all three. Reports alone are not enough. Backups alone are not readable without Handisoft.", bold=True)

    add_table(
        doc,
        ["Priority", "Action", "Why"],
        [
            ["1", "Export CSV/Excel reports while software still runs", "Easiest path into LodgeiT, Xero, spreadsheets"],
            ["2", "Copy entire \\HSoft folder", "Programs, ODB .dat files, documents, config"],
            ["3", "SQL .bak backup (if SQL edition)", "Full database snapshot"],
            ["4", "Copy HSoft\\Doc\\DocBase + W:\\ABC_Data", "Client PDFs, letters, scans, templates"],
            ["5", "Obtain client consent before transferring TFNs/files", "TPB Code + Privacy Act requirement"],
        ],
    )

    # Section 2
    add_heading(doc, "2. Before You Start", 1)

    add_heading(doc, "2.1 Identify ODB vs SQL Edition", 2)
    add_para(doc, "Open any Handisoft program → Help → About:")
    add_table(
        doc,
        ["Indicator", "ODB (Flat File)", "SQL Edition"],
        [
            ["Icon colour", "White icon", "Black icon"],
            ["Help → About", "No \"SQL Link\" shown", "Shows \"SQL Link\""],
            ["Executable names", "Ht24.exe style", "HtSQL24.exe style"],
            ["Primary data store", "HSoft\\Data\\*.dat files", "Microsoft SQL Server database"],
        ],
    )

    add_heading(doc, "2.2 Find Your Server Paths", 2)
    add_para(doc, "Common Handisoft paths (drive letter varies by firm):")
    add_bullet(doc, "W:\\HSoft\\ or \\\\YourServer\\HSoft\\")
    add_bullet(doc, "HSoft\\Doc\\DocBase\\ — client documents")
    add_bullet(doc, "W:\\ABC_Data\\ — may be MYOB Accountants Office naming; copy if it exists")
    add_para(
        doc,
        "Note: W:\\ABC_Data is often a MYOB document store, not standard Handisoft naming. "
        "Many firms map the server as W:. Copy BOTH HSoft and ABC_Data if both exist.",
        italic=True,
    )

    add_heading(doc, "2.3 Get Everyone Out of Handisoft", 2)
    add_para(doc, "Before any backup: ALL staff must close Handisoft programs.")
    add_bullet(doc, "Server: Computer Management → Shared Folders → Open Files")
    add_bullet(doc, "Close any open HSoft files before copying")

    # Section 3
    add_heading(doc, "3. Where Your Data Lives", 1)

    add_heading(doc, "3.1 HSoft Folder Structure", 2)
    structure = (
        "HSoft\\\n"
        "├── Apps\\          ← Programs (HtSQL24.exe, HLSQL.exe, HsSQL.ini)\n"
        "├── Data\\          ← ODB databases (.dat files)\n"
        "├── Doc\\\n"
        "│   └── DocBase\\   ← Document Manager files (COPY THIS)\n"
        "└── (other module folders)"
    )
    p = doc.add_paragraph()
    run = p.add_run(structure)
    run.font.name = "Courier New"
    run.font.size = Pt(9)

    add_table(
        doc,
        ["Component", "ODB Edition", "SQL Edition"],
        [
            ["Client master (CRM)", "HSoft\\Data\\ .dat files", "SQL Server DB"],
            ["HandiTax returns", "Per-year app + data", "SQL DB"],
            ["Practice Manager Contact", "Shared client DB", "SQL DB"],
            ["Documents (PDFs/scans)", "HSoft\\Doc\\DocBase (filesystem)", "Same — not only in SQL"],
            ["ATO/SBR credentials", "keystore.xml", "Same"],
        ],
    )

    # Section 4
    add_heading(doc, "4. Practice Manager Contact (CRM) Export", 1)
    add_para(doc, "Practice Manager Contact is your main CRM module: clients, contacts, groups, diary, file notes.")

    add_heading(doc, "4.1 Current Client List", 2)
    add_numbered(doc, "Open Practice Manager Contact")
    add_numbered(doc, "Reports → Clients → Client List Abbreviated")
    add_numbered(doc, "Filter: Current@Clients = True (or blank for all clients)")
    add_numbered(doc, "Click OK")
    add_numbered(doc, "File → Export to Excel → Generate")
    add_numbered(doc, "Save as: Contact_CurrentClients_YYYY-MM-DD.xlsx")

    add_heading(doc, "4.2 Client Email Report", 2)
    add_numbered(doc, "Reports → Edit Reports")
    add_numbered(doc, "Copy Clients List Abbreviated → rename (e.g. Client Emails)")
    add_numbered(doc, "In report designer: replace TFN column with Email@Clients")
    add_numbered(doc, "Report → Execute → File → Export to Excel")

    add_heading(doc, "4.3 Client Groups Report", 2)
    add_numbered(doc, "Reports → Edit Reports → copy Clients List Abbreviated")
    add_numbered(doc, "Add columns: GroupRef, GroupName")
    add_numbered(doc, "Filter: All Clients or Is in a Client Group")
    add_numbered(doc, "Save, run, and export to Excel")

    add_heading(doc, "4.4 File Notes", 2)
    add_para(
        doc,
        "There is no one-click \"export all file notes\" option. Options: (1) Full server backup "
        "preserves notes for future Handisoft access; (2) Per-client manual export from client history; "
        "(3) Check Outlook if Contact was synced.",
    )

    # Section 5
    add_heading(doc, "5. HandiTax Client Export", 1)
    add_para(doc, "Repeat for EACH tax year installed (2024, 2023, 2022, etc.).", bold=True)

    add_heading(doc, "5.1 Client List Detailed (Recommended)", 2)
    add_numbered(doc, "Open HandiTax (select the year)")
    add_numbered(doc, "Reports → Client → Client list detailed (or Extra Detailed)")
    add_numbered(doc, "Filter: #All (not just current/tagged subset)")
    add_numbered(doc, "Click OK")
    add_numbered(doc, "File → Export To → Excel")
    add_numbered(doc, "Save as: HandiTax_2024_ClientList_Detailed.xlsx")

    add_heading(doc, "5.2 LodgeiT-Style Export", 2)
    add_bullet(doc, "Choose Client list Detailed or Client list Extra Detailed")
    add_bullet(doc, "Select every column you might need on the export page")
    add_bullet(doc, "System generates CSV automatically")

    add_heading(doc, "5.3 Advanced: ASCII Export", 2)
    add_para(doc, "Tools → Export an ASCII file — for migration to non-Handisoft systems.")

    add_heading(doc, "5.4 Full Client + Return Packages", 2)
    add_numbered(doc, "Press F6 → untag all clients")
    add_numbered(doc, "Tag required clients (or all)")
    add_numbered(doc, "Tools → Copy tagged clients to disk")
    add_numbered(doc, "Produces Clients.XXd and Returns.XXd files")
    add_numbered(doc, "Restore later via Tools → Merge Files from Disk")

    # Section 6
    add_heading(doc, "6. Jobs, WIP & Billing", 1)
    add_table(
        doc,
        ["Module", "What to Export", "How"],
        [
            ["Jobflow Manager", "Open jobs, tasks, staff", "Reports → Current Task Staff + Current Task Name → export"],
            ["Time+Billing", "WIP, debtors, uninvoiced", "WIP reports (Aged WIP, Not Invoiced WIP) → Export to Excel"],
        ],
    )

    # Section 7
    add_heading(doc, "7. Full Server Backup Procedure", 1)

    add_heading(doc, "7.1 Copy Entire HSoft Folder", 2)
    p = doc.add_paragraph()
    run = p.add_run(
        'robocopy "\\\\YourServer\\HSoft" "D:\\Handover_Backup\\HSoft" /MIR /R:3 /W:5 /LOG:D:\\Handover_Backup\\robocopy.log'
    )
    run.font.name = "Courier New"
    run.font.size = Pt(9)

    add_heading(doc, "7.2 SQL Database Backup (SQL Edition Only)", 2)
    add_numbered(doc, "Open SQL Server Management Studio (SSMS)")
    add_numbered(doc, "Connect to the SQL instance")
    add_numbered(doc, "Databases → right-click HandiSoft DB → Tasks → Back Up…")
    add_numbered(doc, "Type: Full")
    add_numbered(doc, "Save .bak file to handover drive with date in filename")

    add_heading(doc, "7.3 Copy Document Folders", 2)
    add_bullet(doc, "Copy entire HSoft\\Doc\\DocBase\\ folder")
    add_bullet(doc, "Copy entire W:\\ABC_Data\\ if it exists")
    add_bullet(doc, "Do NOT rename the DocBase root folder")

    add_heading(doc, "7.4 Verify Your Backup", 2)
    add_numbered(doc, "Spot-check 5–10 clients in exported Excel files")
    add_numbered(doc, "Open 5 random client PDFs from DocBase backup")
    add_numbered(doc, "Note backup date, paths, file sizes, and who performed backup")
    add_numbered(doc, "Store one copy off-site")

    doc.add_page_break()

    # Section 8
    add_heading(doc, "8. Searching W:\\ABC_Data for Manuals", 1)
    add_para(doc, "Run these PowerShell commands on a PC with W: drive mapped:")
    p = doc.add_paragraph()
    ps1 = (
        'Get-ChildItem "W:\\ABC_Data" -Recurse -Include *.pdf,*.doc,*.docx `\n'
        '  -ErrorAction SilentlyContinue |\n'
        "  Where-Object { $_.Name -match 'handi|handitax|export|backup|manual' } |\n"
        "  Select-Object FullName, LastWriteTime"
    )
    run = p.add_run(ps1)
    run.font.name = "Courier New"
    run.font.size = Pt(8)

    add_para(doc, "Also search: W:\\ABC_Data\\00 Document Templates")
    add_para(doc, "Look for: HandiTax User Guide, Practice Manager guides, internal backup procedures, report templates.")

    # Section 9
    add_heading(doc, "9. Known Gaps & Workarounds", 1)
    add_table(
        doc,
        ["Data", "In Standard Export?", "Workaround"],
        [
            ["Name, address, phone, email, TFN, ABN", "Yes", "Use Client list detailed"],
            ["Bank account details", "Usually NO", "Check engagement letters / DocBase files; manual entry"],
            ["File notes / interaction history", "No bulk export", "Full backup + per-client if urgent"],
            ["Scanned PDFs / letters", "Not in reports", "Copy DocBase + ABC_Data"],
            ["ATO credentials (SBR)", "Separate", "keystore.xml — only if legally authorised"],
        ],
    )

    # Section 10
    add_heading(doc, "10. Troubleshooting", 1)

    add_heading(doc, "Excel Export: \"Call was rejected by Callee\"", 2)
    add_numbered(doc, "Close ALL Excel windows")
    add_numbered(doc, "Excel → File → Options → General → set default sheets to 3 or more")
    add_numbered(doc, "Disable Excel add-ins")
    add_numbered(doc, "Export one report at a time")
    add_numbered(doc, "Save report to disk WITHOUT \"Open in Excel\", then open manually")
    add_numbered(doc, "Last resort: export to PDF instead")

    add_heading(doc, "Other Common Issues", 2)
    add_table(
        doc,
        ["Error", "Likely Cause", "Fix"],
        [
            ["Cannot read / corrupted .dat", "Network share issue", "Tools → Rebuild Database Files; use Windows SMB"],
            ["Another user changed this record", "Concurrent users", "Get everyone out of Handisoft"],
            ["Missing clients in export", "Wrong filter", "Use #All or Current@Clients = True"],
            ["Blank report", "Corrupt template", "Copy from abbreviated template again"],
        ],
    )

    doc.add_page_break()

    # Section 11
    add_heading(doc, "11. Legal & Privacy Checklist (Australia)", 1)
    add_para(
        doc,
        "Buying the business does NOT automatically give permission to transfer every client file. "
        "Complete this checklist before transferring data.",
    )

    add_checkbox_section(
        doc,
        "Pre-Transfer Legal",
        [
            "Client consent obtained to transfer their file (TPB Code item 6 — confidentiality)",
            "Written notice sent to clients about ownership change",
            "Staff confidentiality agreements in place for sale process",
            "ATO whole-of-practice transfer paperwork completed (if applicable)",
            "Register prepared: what was transferred, when, to whom",
        ],
    )

    add_checkbox_section(
        doc,
        "Data Security",
        [
            "Handover media encrypted (USB / cloud)",
            "TFNs NOT sent via plain email",
            "Access controls and audit trail on handover media",
            "Secure wipe planned for handover media after buyer confirms receipt",
            "Records retained per tax law (generally 5 years from lodgement)",
        ],
    )

    add_para(doc, "References:", bold=True)
    add_bullet(doc, "TPB Confidentiality: https://www.tpb.gov.au/confidentiality-client-information")
    add_bullet(doc, "TPB TFN Protection: https://www.tpb.gov.au/protecting-your-clients-tfns")
    add_bullet(doc, "OAIC TFN Rule: https://www.oaic.gov.au/privacy/privacy-guidance-for-organisations-and-government-agencies/handling-personal-information/the-privacy-tax-file-number-rule-2015-and-the-protection-of-tax-file-number-information")

    doc.add_page_break()

    # Section 12 - Run sheet
    add_heading(doc, "12. 48-Hour Run Sheet", 1)

    add_heading(doc, "Day 1 — Human-Readable Exports (software must still run)", 2)
    day1_items = [
        "Practice Manager Contact → Current client list → Excel",
        "Practice Manager Contact → Email report → Excel",
        "Practice Manager Contact → Groups report → Excel",
        "HandiTax [YEAR] → Client list detailed → Excel (repeat per year)",
        "Time+Billing → WIP + debtors reports → Excel",
        "Jobflow → Current tasks report → Excel",
        "Spot-check: exported row count ≈ client count in software",
    ]
    add_checkbox_section(doc, "Day 1 Checklist", day1_items)

    add_heading(doc, "Day 2 — Machine-Readable Backups (ALL users out of Handisoft)", 2)
    day2_items = [
        "SQL full .bak backup completed (if SQL edition)",
        "Entire HSoft folder copied to encrypted external drive",
        "HSoft\\Doc\\DocBase copied (file size verified)",
        "W:\\ABC_Data copied (templates + legacy docs)",
        "Backup date, paths, and performer documented",
        "One copy stored off-site",
        "Test: 5 random client PDFs open correctly from backup",
    ]
    add_checkbox_section(doc, "Day 2 Checklist", day2_items)

    # Section 13
    add_heading(doc, "13. Useful Links & Contacts", 1)
    add_table(
        doc,
        ["Topic", "URL"],
        [
            ["All technical / backup guides", "https://help-handisoft.theaccessgroup.com/en/collections/15031398-handisoft-technical"],
            ["Migrate to new server", "https://help-handisoft.theaccessgroup.com/en/articles/12129336-handisoft-migrate-your-handisoft-programs-to-a-new-server"],
            ["ODB vs SQL", "https://help-handisoft.theaccessgroup.com/en/articles/12145229-handisoft-odb-or-sql-version"],
            ["Practice Manager Contact", "https://help-handisoft.theaccessgroup.com/en/collections/15173332-practice-manager-contact"],
            ["HandiTax reports", "https://help-handisoft.theaccessgroup.com/en/collections/16275544-handitax-reports-and-the-report-builder"],
            ["Document Manager", "https://help-handisoft.theaccessgroup.com/en/collections/14962730-document-manager"],
            ["LodgeiT migration guide", "https://help.lodgeit.net.au/support/solutions/articles/60000609299"],
            ["Excel export error fix", "https://help-handisoft.theaccessgroup.com/en/articles/13346537-handisoft-call-rejected-by-callee-error-when-exporting-reports-to-excel"],
        ],
    )
    add_para(doc, "Access Group Handisoft Support: 1800 660 670", bold=True)

    doc.add_page_break()

    # Section 14 - Sign off
    add_heading(doc, "14. Sign-Off Page", 1)
    add_para(doc, "Complete this section when handover is finished.")

    add_table(
        doc,
        ["Field", "Details"],
        [
            ["Practice name (seller)", ""],
            ["Practice name (buyer)", ""],
            ["Handover date", ""],
            ["Server path(s)", ""],
            ["Edition (ODB / SQL)", ""],
            ["Tax years exported", ""],
            ["Backup location(s)", ""],
            ["Total clients exported", ""],
            ["Performed by", ""],
            ["Verified by", ""],
        ],
    )

    add_para(doc, "")
    add_para(doc, "Seller representative signature: _________________________________    Date: ______________")
    add_para(doc, "")
    add_para(doc, "Buyer representative signature: _________________________________    Date: ______________")
    add_para(doc, "")
    add_para(doc, "Notes / exceptions:")
    doc.add_paragraph("_" * 80)
    doc.add_paragraph("_" * 80)
    doc.add_paragraph("_" * 80)

    # Footer note
    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f_run = footer.add_run(
        "This document is a practical guide compiled from public Handisoft Help Centre documentation, "
        "third-party migration guides, and industry practice. It is not legal or tax advice. "
        "Consult your lawyer and TPB obligations before transferring client data."
    )
    f_run.font.size = Pt(8)
    f_run.italic = True
    f_run.font.color.rgb = RGBColor(128, 128, 128)

    return doc


def main() -> None:
    output_path = "/workspace/Handisoft_CRM_Data_Handover_Guide.docx"
    doc = build_document()
    doc.save(output_path)
    print(f"Created: {output_path}")


if __name__ == "__main__":
    main()
